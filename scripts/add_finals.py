"""One-off: append finals results to the Entries sheet and create the Finals sheet."""
import uuid, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
P = "data/Tomstars_House_Cup_2026_gsheet.xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["Entries"]
now = datetime.datetime(2026, 10, 27, 13, 0, 0)
HDR = ["sport", "category", "status", "date", "time", "venue", "player1", "house1", "player2", "house2",
       "winner", "winnerHouse", "runnerUp", "runnerUpHouse"]
F = [  # sport, category, status, date, time, venue, p1, h1, p2, h2, winner, wh, runner, rh
 ("Darts","Men","TBA","","","","","","","","","","",""),
 ("Darts","Women","TBA","","","","","","","","","","",""),
 ("Jenga","Men","Completed","Today","","Game Room","","","","","Yogesh Nimbalkar","Blue","Nitin Patil","Green"),
 ("Jenga","Women","Scheduled","Tomorrow","","Game Room","Sneha Deshpande","White","Sapana Husare","White","","","",""),
 ("Chess","Men","TBA","","","","","","","","","","",""),
 ("Chess","Women","TBA","","","","","","","","","","",""),
 ("Carrom","Men","Scheduled","Today","4:30 PM","Majauli","","","","","","","",""),
 ("Carrom","Women","Completed","Today","","","","","","","Ankita Agrawal","Blue","Vaishali Patil","Blue"),
 ("Pool","Men","TBA","","","","","","","","","","",""),
 ("Pool","Women","TBA","","","","","","","","","","",""),
 ("Table Tennis","Men","Scheduled","Today","3:00 PM","Game Room","Satyam Kesarwani","White","Ankit Tyagi","Blue","","","",""),
 ("Table Tennis","Women","Scheduled","Today","5:00 PM","Game Room","Shristi Painaik","Green","Aditi Sabnis","Blue","","","",""),
]
if "Finals" not in wb.sheetnames:
    fs = wb.create_sheet("Finals"); fs.append(HDR)
    for r in F: fs.append(list(r))
    for c in fs[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="13131C")
    for col in fs.columns: fs.column_dimensions[col[0].column_letter].width = 18
    fs.freeze_panes = "A2"
    for i, r in enumerate([r for r in F if r[10]]):
        sport, cat = r[0], r[1]
        for k, (name, house, pts, label) in enumerate([(r[10], r[11], 100, "Final Winner"), (r[12], r[13], 50, "Final Runner Up")]):
            ts = now + datetime.timedelta(seconds=i * 2 + k)
            ws.append([str(uuid.uuid4()), house, pts, f"{sport} ({cat}) - {label}", "Olympia", 1, name, None, ts])
wb.save(P)
